"""
documents.py
============
Extracts plain text from uploaded documents so it can be inserted into
the LLM's context window. Supports PDF (pypdf), Word (.docx, via
python-docx) and plain text files - all handled with free, open-source
libraries, entirely offline.
"""

from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader

SUPPORTED_TEXT_EXTENSIONS = {".txt", ".md", ".csv", ".json", ".log"}


class UnsupportedFileTypeError(ValueError):
    pass


def detect_file_type(filename: str) -> str:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        return "pdf"
    if ext == ".docx":
        return "docx"
    if ext in SUPPORTED_TEXT_EXTENSIONS:
        return "text"
    if ext in {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp"}:
        return "image"
    raise UnsupportedFileTypeError(f"Dateityp '{ext}' wird nicht unterstuetzt.")


def extract_text(file_path: Path, file_type: str) -> str:
    """Returns extracted text, truncated to a sane size so it doesn't
    blow the model's context window."""
    if file_type == "pdf":
        text = _extract_pdf(file_path)
    elif file_type == "docx":
        text = _extract_docx(file_path)
    elif file_type == "text":
        text = file_path.read_text(encoding="utf-8", errors="ignore")
    else:
        return ""

    max_chars = 20_000
    if len(text) > max_chars:
        text = text[:max_chars] + "\n\n[... Inhalt gekuerzt ...]"
    return text.strip()


def _extract_pdf(file_path: Path) -> str:
    reader = PdfReader(str(file_path))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n\n".join(pages)


def _extract_docx(file_path: Path) -> str:
    doc = DocxDocument(str(file_path))
    paragraphs = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            paragraphs.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(paragraphs)
